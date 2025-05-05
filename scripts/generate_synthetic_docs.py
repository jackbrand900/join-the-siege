import argparse
import os
import json
import random
import pandas as pd
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas
from pdf2image import convert_from_path
from docx import Document
from openpyxl import Workbook
from faker import Faker

# Constants for paths and defaults
OUTPUT_DIR = "files/synthetic"
LABELS_PATH = "files/labels.csv"
TEMPLATE_DIR = "templates"
DEFAULT_NUM = 10
ALLOWED_EXTENSIONS = ["pdf", "jpg", "png", "docx", "xlsx"]

# Initialize Faker
fake = Faker()

# Generate realistic dummy values for various fields (would make this modular for production, hardcoding options for now)
def fake_value(field):
    field = field.lower()
    if "name" in field:
        return fake.name()
    elif "id" in field or "ssn" in field or "account" in field:
        return str(random.randint(100000000, 999999999))
    elif "email" in field:
        return fake.email()
    elif "phone" in field or "mobile" in field or "contact" in field:
        return fake.phone_number()
    elif "address" in field:
        return fake.address().replace("\n", ", ")
    elif "date" in field or "dob" in field:
        return fake.date_between(start_date="-5y", end_date="today").strftime("%m/%d/%Y")
    elif "amount" in field or "salary" in field or "total" in field or "payment" in field or "price" in field:
        return f"${round(random.uniform(100, 10000), 2):,.2f}"
    elif "city" in field:
        return fake.city()
    elif "state" in field:
        return fake.state()
    elif "country" in field:
        return fake.country()
    elif "zip" in field or "postal" in field:
        return fake.postcode()
    elif "company" in field or "employer" in field:
        return fake.company()
    elif "job" in field or "position" in field or "title" in field:
        return fake.job()
    elif "bank" in field:
        return random.choice(["Chase", "Bank of America", "Wells Fargo", "Citibank", "HSBC"])
    elif "currency" in field:
        return random.choice(["USD", "EUR", "GBP", "JPY", "CAD"])
    else:
        return random.choice([
            str(random.randint(1000, 9999)),
            fake.word().capitalize(),
            fake.bothify(text='??##')
        ])

# Generate a PDF file with given content
def generate_pdf(content: str, pdf_path: str):
    c = canvas.Canvas(pdf_path, pagesize=LETTER)
    for i, line in enumerate(content.strip().split("\n")):
        c.drawString(100, 750 - 15 * i, line)
    c.save()

# Convert the first page of a PDF to an image format (JPG or PNG)
def generate_image_from_pdf(pdf_path, image_path, fmt):
    pages = convert_from_path(pdf_path, first_page=1, last_page=1)
    format_map = {
        "jpg": "JPEG",
        "jpeg": "JPEG",
        "png": "PNG"
    }
    save_format = format_map.get(fmt.lower(), fmt.upper())
    pages[0].save(image_path, save_format)


# Generate a DOCX file with given content
def generate_docx(content: str, docx_path: str):
    doc = Document()
    for line in content.strip().split("\n"):
        doc.add_paragraph(line)
    doc.save(docx_path)

# Generate an XLSX file with given content
def generate_xlsx(content: str, xlsx_path: str):
    wb = Workbook()
    ws = wb.active
    for i, line in enumerate(content.strip().split("\n"), 1):
        ws.cell(row=i, column=1, value=line)
    wb.save(xlsx_path)

# Construct document content from template fields
def build_template(label: str, fields: list[str]) -> str:
    lines = [f"Document Type: {label.upper()}"]
    for field in fields:
        value = fake_value(field)
        lines.append(f"{field}: {value}")
    return "\n".join(lines)

# Generate documents based on an existing template
def generate_docs(label: str, num: int):
    label = label.lower()
    template_path = os.path.join(TEMPLATE_DIR, f"{label}.json")
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"No template found for label '{label}'")

    with open(template_path, "r") as f:
        template = json.load(f)
        fields = [f["label"] for f in template["fields"]]

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    new_rows = []

    for i in range(num):
        content = build_template(label, fields)
        base = f"{label}_synth_{i}"
        chosen_ext = random.choice(ALLOWED_EXTENSIONS)

        if chosen_ext == "pdf":
            pdf_path = os.path.join(OUTPUT_DIR, f"{base}.pdf")
            generate_pdf(content, pdf_path)
            new_rows.append({"filename": f"{base}.pdf", "label": label})

        elif chosen_ext in {"jpg", "png"}:
            temp_pdf = os.path.join(OUTPUT_DIR, f"{base}_temp.pdf")
            image_path = os.path.join(OUTPUT_DIR, f"{base}.{chosen_ext}")
            generate_pdf(content, temp_pdf)
            generate_image_from_pdf(temp_pdf, image_path, chosen_ext)
            os.remove(temp_pdf)
            new_rows.append({"filename": f"{base}.{chosen_ext}", "label": label})

        elif chosen_ext == "docx":
            docx_path = os.path.join(OUTPUT_DIR, f"{base}.docx")
            generate_docx(content, docx_path)
            new_rows.append({"filename": f"{base}.docx", "label": label})

        elif chosen_ext == "xlsx":
            xlsx_path = os.path.join(OUTPUT_DIR, f"{base}.xlsx")
            generate_xlsx(content, xlsx_path)
            new_rows.append({"filename": f"{base}.xlsx", "label": label})

    df_new = pd.DataFrame(new_rows)

    if os.path.exists(LABELS_PATH):
        df_existing = pd.read_csv(LABELS_PATH)
        df = pd.concat([df_existing, df_new], ignore_index=True).drop_duplicates()
    else:
        df = df_new

    df.to_csv(LABELS_PATH, index=False)
    print(f"Generated {len(new_rows)} new files for label '{label}'")

# Command-line interface
if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--label", required=True, help="Label to generate")
    parser.add_argument("--num", type=int, default=DEFAULT_NUM, help="Number of samples to generate")
    args = parser.parse_args()

    generate_docs(args.label, args.num)